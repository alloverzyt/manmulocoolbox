import os
from typing import Any, Dict, List

from ...core.plugin_interface import BasePlugin, PluginInput, PluginResult
from ...utils.file_utils import ensure_dir, get_output_dir, get_filename


class MarkdownConvertPlugin(BasePlugin):
    @property
    def id(self) -> str:
        return "markdown_convert"

    @property
    def name(self) -> str:
        return "Markdown转换"

    @property
    def description(self) -> str:
        return "将Markdown文档转换为HTML或Word格式"

    @property
    def category(self) -> str:
        return "document"

    @property
    def icon(self) -> str:
        return "edit"

    @property
    def version(self) -> str:
        return "1.0.0"

    @property
    def supported_input_formats(self) -> List[str]:
        return ["md", "markdown"]

    @property
    def supported_output_formats(self) -> List[str]:
        return ["html", "docx"]

    @property
    def options_schema(self) -> Dict[str, Any]:
        return {
            "target": {
                "type": "select",
                "label": "目标格式",
                "default": "html",
                "options": [
                    {"label": "HTML (网页)", "value": "html"},
                    {"label": "Word文档", "value": "docx"}
                ]
            },
            "embed_style": {
                "type": "boolean",
                "label": "内嵌CSS样式",
                "default": True
            }
        }

    def execute(self, input_data: PluginInput, progress_callback=None) -> PluginResult:
        file_path = input_data.file_paths[0]
        options = input_data.options
        target = options.get("target", "html")
        embed_style = options.get("embed_style", True)

        output_dir = get_output_dir(file_path, options.get("output_dir", ""))
        ensure_dir(output_dir)

        base_name = get_filename(file_path)
        ext = "html" if target == "html" else "docx"
        output_path = os.path.join(output_dir, f"{base_name}.{ext}")

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                md_content = f.read()

            if target == "html":
                import markdown
                html_body = markdown.markdown(
                    md_content,
                    extensions=['tables', 'fenced_code', 'codehilite']
                )

                if embed_style:
                    full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
body {{ font-family: "Microsoft YaHei", Arial, sans-serif; max-width: 800px; margin: 20px auto; padding: 0 20px; line-height: 1.8; color: #333; }}
h1, h2, h3 {{ color: #1a1a2e; }}
h1 {{ border-bottom: 2px solid #1a1a2e; padding-bottom: 10px; }}
table {{ border-collapse: collapse; width: 100%; margin: 15px 0; }}
th, td {{ border: 1px solid #ddd; padding: 8px 12px; text-align: left; }}
th {{ background-color: #f5f5f5; }}
code {{ background-color: #f5f5f5; padding: 2px 6px; border-radius: 3px; font-family: Consolas, monospace; }}
pre {{ background-color: #f5f5f5; padding: 15px; border-radius: 5px; overflow-x: auto; }}
pre code {{ background: none; padding: 0; }}
blockquote {{ border-left: 4px solid #ddd; margin: 15px 0; padding: 5px 15px; color: #666; }}
</style>
</head>
<body>
{html_body}
</body>
</html>"""
                else:
                    full_html = f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8"><title>{base_name}</title></head>
<body>{html_body}</body>
</html>"""

                with open(output_path, 'w', encoding='utf-8') as f:
                    f.write(full_html)

            elif target == "docx":
                try:
                    from docx import Document
                except ImportError:
                    return PluginResult(success=False, error="python-docx 未安装。请运行: pip install python-docx")

                doc = Document()

                import markdown
                html_body = markdown.markdown(md_content, extensions=['tables', 'fenced_code'])

                from docx.shared import Pt, Inches
                from docx.enum.text import WD_ALIGN_PARAGRAPH

                for line in md_content.split('\n'):
                    stripped = line.strip()
                    if not stripped:
                        continue

                    if stripped.startswith('# '):
                        doc.add_heading(stripped[2:], level=1)
                    elif stripped.startswith('## '):
                        doc.add_heading(stripped[3:], level=2)
                    elif stripped.startswith('### '):
                        doc.add_heading(stripped[4:], level=3)
                    elif stripped.startswith('- ') or stripped.startswith('* '):
                        doc.add_paragraph(stripped[2:], style='List Bullet')
                    elif stripped.startswith('> '):
                        p = doc.add_paragraph(stripped[2:])
                        p.style = doc.styles['Quote'] if 'Quote' in [s.name for s in doc.styles] else p.style
                    else:
                        clean = stripped
                        clean = clean.replace('**', '').replace('__', '')
                        clean = clean.replace('*', '').replace('_', '')
                        clean = clean.replace('`', '')
                        doc.add_paragraph(clean)

                doc.save(output_path)

            return PluginResult(
                success=True,
                output_paths=[output_path],
                message=f"Markdown转{target.upper()}成功"
            )

        except ImportError as e:
            return PluginResult(success=False, error=f"缺少依赖库: {str(e)}")
        except Exception as e:
            return PluginResult(success=False, error=f"转换失败: {str(e)}")
